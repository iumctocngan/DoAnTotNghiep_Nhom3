from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from generation.schemas import GeneratedLearningOutcome

BORDER = Border(*(Side(style="thin", color="D1D5DB") for _ in range(4)))

LEVEL_NAMES = (
    "1: Chưa đạt",
    "2: Cần cải thiện",
    "3: Đạt",
    "4: Tốt",
    "5: Rất tốt",
)

META_NAMES = (
    (5, "Mã CLO", 24),
    (6, "Chuẩn đầu ra (CLO)", 45),
    (7, "Tiêu chí đánh giá", 65),
)


def style_cell(cell, font_args, align_args=None):
    cell.font = Font(name="Arial", **font_args)
    cell.alignment = Alignment(**(align_args or {"vertical": "top", "wrap_text": True}))
    cell.border = BORDER


def format_header(title: str, desc: str) -> str:
    t, d = title.strip().lstrip("-").strip(), desc.strip().lstrip("-").strip()
    if not d or d == t:
        return f"- {t}."
    if t.lower() in d.lower():
        return f"- {d}"
    return f"- {t}:\n{d}"


def render_xlsx(outcomes: list[GeneratedLearningOutcome]) -> bytes:
    wb = Workbook()
    s1 = wb.active
    s1.title = "Ma trận Rubric"
    s1.column_dimensions["A"].width = 22

    s1["A1"] = "Chuẩn đầu ra (CLO) & Ma trận Rubric đánh giá"
    s1["A1"].font = Font(name="Arial", size=14, bold=True)
    s1["A2"] = (
        "1: Chưa đạt · 2: Cần cải thiện · 3: Đạt · 4: Tốt · 5: Rất tốt. "
        "Đánh giá mức độ đạt được của học sinh theo từng tiêu chí năng lực trọng tâm."
    )
    s1["A2"].font = Font(name="Arial", size=9.5, italic=True)
    for r, h in ((1, 26), (2, 20), (3, 10), (4, 10)):
        s1.row_dimensions[r].height = h

    for r, label, h in META_NAMES:
        s1.row_dimensions[r].height = h
        style_cell(s1.cell(r, 1, label), {"size": 10, "bold": True}, {"vertical": "center", "wrap_text": True})

    for i, label in enumerate(LEVEL_NAMES):
        r = 8 + i
        s1.row_dimensions[r].height = 58
        style_cell(s1.cell(r, 1, label), {"size": 10, "bold": True}, {"horizontal": "center", "vertical": "center", "wrap_text": True})

    col = 2
    for o in outcomes:
        for c in o.criteria:
            s1.column_dimensions[get_column_letter(col)].width = 38
            style_cell(s1.cell(5, col, o.code), {"size": 10, "bold": True}, {"horizontal": "center", "vertical": "center"})
            style_cell(s1.cell(6, col, o.description), {"size": 9.5})
            style_cell(s1.cell(7, col, format_header(c.title, c.description)), {"size": 10, "bold": True})
            for i in range(5):
                style_cell(s1.cell(8 + i, col, c.levels[i].description if i < len(c.levels) else ""), {"size": 9.5})
            col += 1

    s2 = wb.create_sheet(title="Danh sách CLO chi tiết")
    s2.append(("STT", "Mã CLO", "Chuẩn đầu ra", "Tiêu chí đánh giá", *LEVEL_NAMES))
    s2.row_dimensions[1].height = 26
    for idx, w in enumerate((6, 12, 34, 34, 28, 28, 28, 28, 28), start=1):
        s2.column_dimensions[get_column_letter(idx)].width = w
        style_cell(s2.cell(1, idx), {"size": 10, "bold": True}, {"horizontal": "center", "vertical": "center", "wrap_text": True})

    stt = 1
    for o in outcomes:
        for c in o.criteria:
            s2.append((stt, o.code, o.description, format_header(c.title, c.description), *(c.levels[i].description if i < len(c.levels) else "" for i in range(5))))
            r = s2.max_row
            s2.row_dimensions[r].height = 58
            for c_idx in range(1, 10):
                align = {"horizontal": "center", "vertical": "top"} if c_idx in (1, 2) else None
                style_cell(s2.cell(r, c_idx), {"size": 9.5}, align_args=align)
            stt += 1
    s2.auto_filter.ref = s2.dimensions

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
